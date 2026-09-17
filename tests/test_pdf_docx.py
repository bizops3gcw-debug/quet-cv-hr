import io
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pymupdf as fitz
import docx
from core.document_loader import DocumentLoader


def test_create_and_load_pdf():
    """Tạo một file PDF kỹ thuật số mẫu bằng PyMuPDF và nạp qua DocumentLoader."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "CV - NGUYEN VAN A\nVi tri: Python Backend Developer\nSDT: 0987654321", fontsize=12)
    pdf_bytes = doc.write()
    doc.close()

    loader = DocumentLoader()
    images, text = loader.load_document(pdf_bytes, "test_resume.pdf")

    assert len(images) == 1
    assert "Python Backend Developer" in text
    print(" Kiểm tra nạp file PDF thành công!")


def test_create_and_load_docx():
    """Tạo một file Word DOCX mẫu bằng python-docx và nạp qua DocumentLoader."""
    doc = docx.Document()
    doc.add_heading("RESUME - JOHN DOE", level=1)
    doc.add_paragraph("Target: Senior Cloud Engineer")
    doc.add_paragraph("Email: john.doe@example.com")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "AWS, Kubernetes, Terraform"
    table.cell(1, 0).text = "Experience"
    table.cell(1, 1).text = "5 years"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    loader = DocumentLoader()
    images, text = loader.load_document(docx_bytes, "test_resume.docx")

    assert "Senior Cloud Engineer" in text
    assert "AWS, Kubernetes, Terraform" in text
    print(" Kiểm tra nạp file Word DOCX thành công!")


if __name__ == "__main__":
    test_create_and_load_pdf()
    test_create_and_load_docx()
